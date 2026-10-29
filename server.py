"""
工事作業予定表 - バックエンドAPI (FastAPI)
"""
from fastapi import FastAPI, HTTPException, Request, UploadFile, File
from fastapi.staticfiles import StaticFiles
from fastapi.responses import FileResponse, JSONResponse
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel
from typing import Optional, List
import hashlib
import hmac
import io
import json
import time
import os
import re
import uuid
import asyncio
import base64
import urllib.request
import urllib.error
import threading
from datetime import datetime, date, timezone, timedelta
import calendar
try:
    import jpholiday
    HAS_JPHOLIDAY = True
except ImportError:
    HAS_JPHOLIDAY = False

try:
    import holidays as _holidays_lib
    HAS_HOLIDAYS = True
except ImportError:
    HAS_HOLIDAYS = False

# 中国祝日名の日本語訳テーブル（部分マッチ）
CN_NAME_JA = {
    "元旦": "中国：元旦",
    "春节": "中国：春節（旧正月）",
    "农历除夕": "中国：大晦日（旧暦）",
    "清明节": "中国：清明節",
    "劳动节": "中国：労働節",
    "端午节": "中国：端午節",
    "中秋节": "中国：中秋節",
    "国庆节": "中国：国慶節",
    "休息日": "中国：振替休日",
}

# 韓国祝日名の日本語訳テーブル
KR_NAME_JA = {
    "신정연휴": "韓国：元旦",
    "설날 전날": "韓国：ソルラル前日",
    "설날": "韓国：ソルラル（旧正月）",
    "설날 다음날": "韓国：ソルラル翌日",
    "삼일절": "韓国：三一節",
    "삼일절 대체 휴일": "韓国：三一節振替",
    "노동절": "韓国：労働節",
    "어린이날": "韓国：こどもの日",
    "부처님오신날": "韓国：仏誕節",
    "부처님오신날 대체 휴일": "韓国：仏誕節振替",
    "현충일": "韓国：顕忠日",
    "광복절": "韓国：光復節",
    "광복절 대체 휴일": "韓国：光復節振替",
    "추석 전날": "韓国：秋夕前日",
    "추석": "韓국：秋夕（チュソク）",
    "추석 다음날": "韓国：秋夕翌日",
    "개천절": "韓国：開天節",
    "개천절 대체 휴일": "韓国：開天節振替",
    "한글날": "韓国：ハングルの日",
    "기독탄신일": "韓国：クリスマス",
    "지방선거일": "韓国：地方選挙日",
    "대통령선거일": "韓国：大統領選挙日",
    "국회의원선거일": "韓国：国会議員選挙日",
}

def _cn_holiday_name_ja(raw: str) -> str:
    for key, ja in CN_NAME_JA.items():
        if key in raw:
            return ja
    return f"中国：{raw}"

def _kr_holiday_name_ja(raw: str) -> str:
    return KR_NAME_JA.get(raw, f"韓国：{raw}")

# 年ごとのキャッシュ
_cn_cache: dict = {}
_kr_cache: dict = {}

def _get_cn_holidays(year: int):
    if year not in _cn_cache:
        _cn_cache[year] = _holidays_lib.China(years=year) if HAS_HOLIDAYS else {}
    return _cn_cache[year]

def _get_kr_holidays(year: int):
    if year not in _kr_cache:
        _kr_cache[year] = _holidays_lib.SouthKorea(years=year) if HAS_HOLIDAYS else {}
    return _kr_cache[year]

app = FastAPI(title="工事作業予定表API")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

# -------------------------------------------------------
# パス設定
# -------------------------------------------------------
_BASE_DIR = os.path.dirname(os.path.abspath(__file__))
# Render無料プランはディスク永続化なし → ローカルはキャッシュ用のみ
DATA_FILE    = os.path.join(_BASE_DIR, "data.json")
INITIAL_FILE = os.path.join(_BASE_DIR, "initial_data.json")

# -------------------------------------------------------
# GitHub API による永続化設定
# 環境変数 GITHUB_TOKEN / GITHUB_REPO / GITHUB_DATA_PATH を設定する
# GITHUB_REPO  例: "aihara1089-creator/monitoring-gr-koji-yotei-"
# GITHUB_DATA_PATH 例: "data.json"（リポジトリルートからのパス）
# -------------------------------------------------------
GITHUB_TOKEN     = os.environ.get("GITHUB_TOKEN", "")
GITHUB_REPO      = os.environ.get("GITHUB_REPO", "aihara1089-creator/monitoring-gr-koji-yotei-")
GITHUB_DATA_PATH = os.environ.get("GITHUB_DATA_PATH", "data.json")
GITHUB_BRANCH    = os.environ.get("GITHUB_BRANCH", "main")
_GH_API_BASE     = f"https://api.github.com/repos/{GITHUB_REPO}/contents/{GITHUB_DATA_PATH}"

# 書き込み排他制御用ロック（API間の競合防止）
_DATA_LOCK = asyncio.Lock()

# GitHub push用スレッドロック（複数スレッドの同時push防止）
_GH_PUSH_LOCK = threading.Lock()

# GitHub上のファイルの現在のSHA（更新時に必要）
_gh_sha_cache: str = ""

def _gh_headers():
    h = {"Accept": "application/vnd.github+json", "X-GitHub-Api-Version": "2022-11-28"}
    if GITHUB_TOKEN:
        h["Authorization"] = f"Bearer {GITHUB_TOKEN}"
    return h

def _gh_fetch() -> dict | None:
    """GitHub から data.json を取得。失敗時は None を返す。"""
    global _gh_sha_cache  # ← 関数先頭で宣言
    if not GITHUB_TOKEN:
        return None
    try:
        req = urllib.request.Request(
            f"{_GH_API_BASE}?ref={GITHUB_BRANCH}",
            headers=_gh_headers()
        )
        with urllib.request.urlopen(req, timeout=10) as resp:
            meta = json.loads(resp.read())
        _gh_sha_cache = meta.get("sha", "")
        raw = base64.b64decode(meta["content"]).decode("utf-8")
        return json.loads(raw)
    except Exception as e:
        print(f"[GitHub fetch] failed: {e}")
        return None

def _gh_push(data: dict):
    """
    ローカルの最新データを GitHub の data.json に push する。
    スレッドロックで直列化し、ロック取得時点のローカルファイルを使う。
    複数リクエストが連続した場合でも、常に最新データがpushされる。
    """
    global _gh_sha_cache
    if not GITHUB_TOKEN:
        return
    with _GH_PUSH_LOCK:  # 同時push防止：1スレッドずつ直列実行
        # ★ ロック取得後に改めてローカルファイルの最新を読み直す
        # （ロック待ち中に別スレッドがより新しいデータをローカル保存している可能性があるため）
        try:
            with open(DATA_FILE, "r", encoding="utf-8") as f:
                latest_data = json.load(f)
        except Exception:
            latest_data = data  # 読み込み失敗時は引数のデータを使う

        for attempt in range(3):  # 最大3回試行（SHA競合・ネットワークエラー対策）
            try:
                content_b64 = base64.b64encode(
                    json.dumps(latest_data, ensure_ascii=False, indent=2).encode("utf-8")
                ).decode("ascii")
                payload = {
                    "message": "chore: auto-save data.json via API",
                    "content": content_b64,
                    "branch": GITHUB_BRANCH,
                }
                if _gh_sha_cache:
                    payload["sha"] = _gh_sha_cache
                body = json.dumps(payload).encode("utf-8")
                req = urllib.request.Request(
                    _GH_API_BASE,
                    data=body,
                    headers={**_gh_headers(), "Content-Type": "application/json"},
                    method="PUT"
                )
                with urllib.request.urlopen(req, timeout=20) as resp:
                    result = json.loads(resp.read())
                _gh_sha_cache = result.get("content", {}).get("sha", _gh_sha_cache)
                print(f"[GitHub push] OK  sha={_gh_sha_cache[:8]}")
                return  # 成功
            except urllib.error.HTTPError as e:
                body_err = e.read().decode("utf-8", errors="replace")
                print(f"[GitHub push] attempt={attempt+1} HTTP {e.code}: {body_err[:200]}")
                if e.code in (409, 422) and attempt < 2:
                    # SHA競合 → 最新SHAを再取得してリトライ
                    print("[GitHub push] SHA conflict → re-fetching SHA")
                    _gh_fetch()  # SHAキャッシュを更新
                    print(f"[GitHub push] SHA refreshed: {_gh_sha_cache[:8]}")
                    continue
                break
            except Exception as e:
                print(f"[GitHub push] attempt={attempt+1} failed: {e}")
                if attempt < 2:
                    import time
                    time.sleep(1)  # 1秒待ってリトライ
                    continue
                break

# -------------------------------------------------------
# データロード/セーブ
# -------------------------------------------------------
def _ensure_fields(d: dict) -> dict:
    """必須フィールドが欠けていれば補完する"""
    if "reminders" not in d:
        d["reminders"] = []
        d["reminder_next_id"] = 1
    # 新造船データ（機密）は公開リポジトリの data.json には入れない。
    # 非公開リポジトリ + パスワード保護の別ストアで管理する（下の「新造船」節を参照）。
    d.pop("vessels", None)
    d.pop("vessels_meta", None)
    return d

def _init_local_data():
    """
    起動時に1回だけ呼ぶ。
    優先順位: GitHub > ローカルキャッシュ > initial_data.json
    結果をローカル data.json に書いてキャッシュとする。
    """
    # ① GitHub から最新を取得
    gh = _gh_fetch()
    if gh is not None:
        _ensure_fields(gh)
        with open(DATA_FILE, "w", encoding="utf-8") as f:
            json.dump(gh, f, ensure_ascii=False, indent=2)
        print(f"[startup] Loaded from GitHub  schedules={len(gh['schedules'])}  staff={len(gh['staff'])}")
        return

    # ② ローカルキャッシュが既にある
    if os.path.exists(DATA_FILE):
        print("[startup] Using local cache data.json (GitHub not configured or unreachable)")
        return

    # ③ 初回: initial_data.json からシード
    if os.path.exists(INITIAL_FILE):
        with open(INITIAL_FILE, "r", encoding="utf-8") as f:
            raw = json.load(f)
        schedules = []
        for i, s in enumerate(raw.get("schedules", [])):
            s["id"] = str(i + 1)
            schedules.append(s)
        data = {
            "schedules": schedules,
            "staff": raw.get("staff", []),
            "next_id": len(schedules) + 1,
            "reminders": [],
            "reminder_next_id": 1,
        }
        with open(DATA_FILE, "w", encoding="utf-8") as f:
            json.dump(data, f, ensure_ascii=False, indent=2)
        print(f"[startup] Seeded from initial_data.json  schedules={len(schedules)}")
        return

    # ④ 完全に空のデータ
    data = {"schedules": [], "staff": [], "next_id": 1, "reminders": [], "reminder_next_id": 1}
    with open(DATA_FILE, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)
    print("[startup] Started with empty data")

# アプリ起動時に実行（インポート時）
_init_local_data()

def load_data() -> dict:
    """ローカルキャッシュから読む（起動後はGitHubにアクセスしない）"""
    if os.path.exists(DATA_FILE):
        with open(DATA_FILE, "r", encoding="utf-8") as f:
            d = json.load(f)
        return _ensure_fields(d)
    return {"schedules": [], "staff": [], "next_id": 1, "reminders": [], "reminder_next_id": 1}

def save_data(data: dict):
    """ローカルに書いた後、GitHub に push する。
    daemon=False でプロセス終了前に必ずスレッドが完了するようにし、
    データ消失（Renderスリープ前にpushが終わらない問題）を防ぐ。
    """
    with open(DATA_FILE, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)
    # daemon=False: プロセス終了時にスレッド完了を待つ（Renderスリープ対策）
    t = threading.Thread(target=_gh_push, args=(data,), daemon=False)
    t.start()

# -------------------------------------------------------
# モデル
# -------------------------------------------------------
class ScheduleCreate(BaseModel):
    year: int
    month: int
    staff: str
    day: int
    content: str
    color: Optional[str] = "#3B82F6"

class ScheduleUpdate(BaseModel):
    content: Optional[str] = None
    color: Optional[str] = None
    staff: Optional[str] = None
    day: Optional[int] = None
    year: Optional[int] = None
    month: Optional[int] = None

class StaffCreate(BaseModel):
    name: str

class StaffOrder(BaseModel):
    staff: List[str]

class ReminderCreate(BaseModel):
    schedule_id: str          # 対象スケジュールID
    remind_days_before: int   # 何日前に通知するか
    note: Optional[str] = ""  # メモ

class ReminderUpdate(BaseModel):
    remind_days_before: Optional[int] = None
    note: Optional[str] = None

# -------------------------------------------------------
# スケジュールAPI
# -------------------------------------------------------
@app.get("/api/schedule")
def get_schedule(year: int, month: int):
    data = load_data()
    entries = [s for s in data["schedules"] if s["year"] == year and s["month"] == month]
    
    # カレンダー情報も返す
    _, days_in_month = calendar.monthrange(year, month)
    cal_days = []
    for d in range(1, days_in_month + 1):
        dt = datetime(year, month, d)
        weekday_names = ["月", "火", "水", "木", "金", "土", "日"]
        d_obj = date(year, month, d)
        # 日本祝日
        is_holiday = HAS_JPHOLIDAY and jpholiday.is_holiday(d_obj)
        holiday_name = (jpholiday.is_holiday_name(d_obj) if HAS_JPHOLIDAY else None) or ""
        # 中国祝日
        cn_hols = _get_cn_holidays(year)
        cn_raw = cn_hols.get(d_obj, "")
        is_cn_holiday = bool(cn_raw)
        cn_holiday_name = _cn_holiday_name_ja(cn_raw) if cn_raw else ""
        # 韓国祝日
        kr_hols = _get_kr_holidays(year)
        kr_raw = kr_hols.get(d_obj, "")
        is_kr_holiday = bool(kr_raw)
        kr_holiday_name = _kr_holiday_name_ja(kr_raw) if kr_raw else ""
        cal_days.append({
            "day": d,
            "weekday": weekday_names[dt.weekday()],
            "is_weekend": dt.weekday() >= 5,
            "is_holiday": bool(is_holiday),
            "holiday_name": holiday_name,
            "is_cn_holiday": is_cn_holiday,
            "cn_holiday_name": cn_holiday_name,
            "is_kr_holiday": is_kr_holiday,
            "kr_holiday_name": kr_holiday_name,
        })
    
    return {
        "year": year,
        "month": month,
        "days": cal_days,
        "schedules": entries,
        "staff": data["staff"]
    }

@app.get("/api/schedules/upcoming")
def get_upcoming_schedules():
    """今日（日本時間）以降の工事予定を返す。船の竣工日とのずれチェック用（予定は公開データのみ）"""
    today = datetime.now(timezone(timedelta(hours=9))).date()
    key = (today.year, today.month, today.day)
    data = load_data()
    out = [
        {"id": s["id"], "year": s["year"], "month": s["month"], "day": s["day"],
         "staff": s.get("staff", ""), "content": s.get("content", "")}
        for s in data["schedules"]
        if (s["year"], s["month"], s["day"]) >= key
    ]
    out.sort(key=lambda s: (s["year"], s["month"], s["day"]))
    resp = JSONResponse({"today": today.isoformat(), "schedules": out})
    resp.headers["Cache-Control"] = "no-store"
    return resp

@app.post("/api/schedule")
async def create_schedule(item: ScheduleCreate):
    async with _DATA_LOCK:
        data = load_data()
        new_id = str(data.get("next_id", 1))
        entry = {
            "id": new_id,
            "year": item.year,
            "month": item.month,
            "staff": item.staff,
            "day": item.day,
            "content": item.content,
            "color": item.color or "#3B82F6"
        }
        data["schedules"].append(entry)
        data["next_id"] = int(new_id) + 1
        save_data(data)
    return entry

@app.put("/api/schedule/{schedule_id}")
async def update_schedule(schedule_id: str, item: ScheduleUpdate):
    async with _DATA_LOCK:
        data = load_data()
        for s in data["schedules"]:
            if s["id"] == schedule_id:
                if item.content is not None:
                    s["content"] = item.content
                if item.color is not None:
                    s["color"] = item.color
                if item.staff is not None:
                    s["staff"] = item.staff
                if item.day is not None:
                    s["day"] = item.day
                if item.year is not None:
                    s["year"] = item.year
                if item.month is not None:
                    s["month"] = item.month
                save_data(data)
                return s
    raise HTTPException(status_code=404, detail="Not found")

@app.delete("/api/schedule/{schedule_id}")
async def delete_schedule(schedule_id: str):
    async with _DATA_LOCK:
        data = load_data()
        original = len(data["schedules"])
        data["schedules"] = [s for s in data["schedules"] if s["id"] != schedule_id]
        if len(data["schedules"]) == original:
            raise HTTPException(status_code=404, detail="Not found")
        save_data(data)
    return {"ok": True}

# -------------------------------------------------------
# 担当者API
# -------------------------------------------------------
@app.get("/api/staff")
def get_staff():
    data = load_data()
    return {"staff": data["staff"]}

@app.post("/api/staff")
async def add_staff(item: StaffCreate):
    async with _DATA_LOCK:
        data = load_data()
        if item.name in data["staff"]:
            raise HTTPException(status_code=400, detail="既に存在します")
        data["staff"].append(item.name)
        save_data(data)
    return {"staff": data["staff"]}

@app.delete("/api/staff/{name}")
async def delete_staff(name: str):
    async with _DATA_LOCK:
        data = load_data()
        if name not in data["staff"]:
            raise HTTPException(status_code=404, detail="Not found")
        data["staff"] = [s for s in data["staff"] if s != name]
        save_data(data)
    return {"staff": data["staff"]}

@app.put("/api/staff/order")
async def reorder_staff(item: StaffOrder):
    async with _DATA_LOCK:
        data = load_data()
        data["staff"] = item.staff
        save_data(data)
    return {"staff": data["staff"]}

# -------------------------------------------------------
# リマインドAPI
# -------------------------------------------------------
@app.get("/api/reminders")
def get_reminders():
    data = load_data()
    reminders = data.get("reminders", [])
    # 今日基準で「あと何日」を計算して付与
    today = date.today()
    result = []
    for r in reminders:
        # 対象スケジュールの日付を探す
        sch = next((s for s in data["schedules"] if s["id"] == r["schedule_id"]), None)
        if not sch:
            continue  # スケジュール削除済みはスキップ
        sch_date = date(sch["year"], sch["month"], sch["day"])
        remind_date = sch_date - __import__('datetime').timedelta(days=r["remind_days_before"])
        days_until_remind = (remind_date - today).days
        days_until_event = (sch_date - today).days
        result.append({
            **r,
            "schedule": sch,
            "remind_date": remind_date.isoformat(),
            "days_until_remind": days_until_remind,
            "days_until_event": days_until_event,
            "is_due": days_until_remind <= 0 <= days_until_event,  # 通知タイミング内
            "is_past": days_until_event < 0,
        })
    return {"reminders": result}

@app.post("/api/reminders")
async def create_reminder(item: ReminderCreate):
    async with _DATA_LOCK:
        data = load_data()
        sch = next((s for s in data["schedules"] if s["id"] == item.schedule_id), None)
        if not sch:
            raise HTTPException(status_code=404, detail="スケジュールが見つかりません")
        new_id = str(data.get("reminder_next_id", 1))
        reminder = {
            "id": new_id,
            "schedule_id": item.schedule_id,
            "remind_days_before": item.remind_days_before,
            "note": item.note or "",
            "created_at": datetime.now().isoformat()
        }
        data.setdefault("reminders", []).append(reminder)
        data["reminder_next_id"] = int(new_id) + 1
        save_data(data)
    return reminder

@app.put("/api/reminders/{reminder_id}")
async def update_reminder(reminder_id: str, item: ReminderUpdate):
    async with _DATA_LOCK:
        data = load_data()
        for r in data.get("reminders", []):
            if r["id"] == reminder_id:
                if item.remind_days_before is not None:
                    r["remind_days_before"] = item.remind_days_before
                if item.note is not None:
                    r["note"] = item.note
                save_data(data)
                return r
    raise HTTPException(status_code=404, detail="Not found")

@app.delete("/api/reminders/{reminder_id}")
async def delete_reminder(reminder_id: str):
    async with _DATA_LOCK:
        data = load_data()
        before = len(data.get("reminders", []))
        data["reminders"] = [r for r in data.get("reminders", []) if r["id"] != reminder_id]
        if len(data["reminders"]) == before:
            raise HTTPException(status_code=404, detail="Not found")
        save_data(data)
    return {"ok": True}

# -------------------------------------------------------
# 月一覧API (存在するデータの年月リスト)
# -------------------------------------------------------
@app.get("/api/months")
def get_months():
    data = load_data()
    months_set = set()
    for s in data["schedules"]:
        months_set.add((s["year"], s["month"]))
    months = sorted(months_set)
    return {"months": [{"year": y, "month": m} for y, m in months]}

# -------------------------------------------------------
# 新造船 竣工予定API（MOL「新造船の工程スケジュール」Excelの取込）
# -------------------------------------------------------
_VESSEL_MAX_BYTES = 10 * 1024 * 1024
_VESSEL_COLS = {
    "BUILDER_YARD": "yard",
    "BUILDERS_VESSEL_NUMBER": "hull_no",
    "VESSEL_NAME": "name",
    "VESSEL_TYPE": "type",
    "OWNERSHIP_TYPE_BEFORE_DELIVERY": "ownership",
    "PLANNED_SEA_TRIALS_DATE": "sea_trial",
    "PLANNED_CONSTRUCTION_COMPLETE_DATE": "complete",
    "PLANNED_DATE_OF_BUILD_DATE": "delivery",
    "変更フラグ": "change_flag",
    "変更内容": "change_note",
}
_YMD_RE = re.compile(r"^\s*(\d{4})[/\-.](\d{1,2})[/\-.](\d{1,2})(?:[ T]\d{1,2}:\d{2}(?::\d{2})?)?\s*$")

def _cell_text(v) -> str:
    """セル値を表示用の文字列にする（空・0 は空文字）"""
    if v is None:
        return ""
    if isinstance(v, datetime):
        return v.strftime("%Y/%m/%d")
    if isinstance(v, date):
        return v.strftime("%Y/%m/%d")
    if isinstance(v, (int, float)):
        if v == 0:
            return ""
        return str(int(v)) if float(v).is_integer() else str(v)
    s = str(v).strip()
    return "" if s in ("", "0") else s

def _to_iso_date(text: str):
    """'2026/12/11' '2027-12-31' のような完全な日付のみ ISO 形式にする。TBD や '2027/03/FH' は None"""
    m = _YMD_RE.match(text or "")
    if not m:
        return None
    try:
        return date(int(m.group(1)), int(m.group(2)), int(m.group(3))).isoformat()
    except ValueError:
        return None

_VESSEL_SOURCES = {"mol": "MOL", "tsuneishi": "常石"}

def _parse_mol(wb):
    """MOL新造船工程表。最新（右端）の、見出しが読めるシートを使う。見つからなければ None"""
    for ws in reversed(wb.worksheets):
        if ws.sheet_state != "visible":
            continue
        rows = [tuple(r) for r in ws.iter_rows(values_only=True)]
        hdr_idx, colmap = None, {}
        for i, row in enumerate(rows[:15]):
            heads = [str(c).strip() if c is not None else "" for c in row]
            if "VESSEL_NAME" in heads and "PLANNED_DATE_OF_BUILD_DATE" in heads:
                hdr_idx = i
                colmap = {_VESSEL_COLS[h]: j for j, h in enumerate(heads) if h in _VESSEL_COLS}
                break
        if hdr_idx is None:
            continue

        # 見出しより上の行から「テーブル名(今月)」の値を拾う
        table_name = ""
        for row in rows[:hdr_idx]:
            for j, c in enumerate(row):
                t = str(c) if c is not None else ""
                if "テーブル名" in t and "今月" in t:
                    for nxt in row[j + 1:]:
                        if nxt not in (None, ""):
                            table_name = str(nxt).strip()
                            break
        vessels, seen = [], {}
        for row in rows[hdr_idx + 1:]:
            def g(field):
                j = colmap.get(field)
                return _cell_text(row[j]) if j is not None and j < len(row) else ""
            name, hull = g("name"), g("hull_no")
            if not name and not hull:
                continue
            yard = g("yard")
            key = f"{yard}|{hull}" if hull else f"{yard}|{name}"
            seen[key] = seen.get(key, 0) + 1
            if seen[key] > 1:
                key = f"{key}#{seen[key]}"
            delivery, complete = g("delivery"), g("complete")
            vessels.append({
                "key": key, "source": "mol", "yard": yard, "hull_no": hull, "work_no": "", "name": name,
                "type": g("type"), "ownership": g("ownership"),
                "sea_trial": g("sea_trial"),
                "complete": complete, "complete_date": _to_iso_date(complete),
                "delivery": delivery, "delivery_date": _to_iso_date(delivery),
                "change_flag": g("change_flag"), "change_note": g("change_note"),
            })
        return vessels, {"sheet": ws.title, "table_name": table_name}
    return None

def _one_line(v) -> str:
    """改行・連続空白を1つの空白にまとめる"""
    return re.sub(r"\s+", " ", _cell_text(v)).strip()

def _parse_tsuneishi(wb):
    """常石造船の標準搭載スケジュール表（工事番号／番船／船名／海上試運転／デリバリー日予定）。
    顧客連絡先などの列は読み込まない。見つからなければ None"""
    sheets = [ws for ws in wb.worksheets if ws.sheet_state == "visible"]
    sheets.sort(key=lambda w: 0 if "常石標準搭載" in w.title else 1)   # 該当名のシートを優先
    for ws in sheets:
        rows = [tuple(r) for r in ws.iter_rows(values_only=True)]
        hdr_idx, colmap = None, {}
        for i, row in enumerate(rows[:15]):
            cm = {}
            for j, c in enumerate(row):
                h = re.sub(r"\s+", "", str(c)) if c is not None else ""
                if h.startswith("工事番号"):
                    cm.setdefault("work_no", j)
                elif h.startswith("番船"):
                    cm.setdefault("hull_no", j)
                elif h.startswith("船名"):
                    cm.setdefault("name", j)
                elif h.startswith("海上試運転"):
                    cm.setdefault("sea_trial", j)
                elif h.startswith("デリバリー日"):
                    cm.setdefault("delivery", j)
            if "hull_no" in cm and "delivery" in cm:
                hdr_idx, colmap = i, cm
                break
        if hdr_idx is None:
            continue
        vessels, seen = [], {}
        for row in rows[hdr_idx + 1:]:
            def g(field):
                j = colmap.get(field)
                return _one_line(row[j]) if j is not None and j < len(row) else ""
            hull, name = g("hull_no"), g("name")
            if name in ("無", "未定", "-", "－"):
                name = ""
            if not hull and not name:
                continue
            key = f"常石|{hull}" if hull else f"常石|{name}"
            seen[key] = seen.get(key, 0) + 1
            if seen[key] > 1:
                key = f"{key}#{seen[key]}"
            delivery = g("delivery")
            vessels.append({
                "key": key, "source": "tsuneishi", "yard": "常石", "hull_no": hull, "work_no": g("work_no"),
                "name": name, "type": "", "ownership": "",
                "sea_trial": g("sea_trial"),
                "complete": "", "complete_date": None,
                "delivery": delivery, "delivery_date": _to_iso_date(delivery),
                "change_flag": "", "change_note": "",
            })
        return vessels, {"sheet": ws.title, "table_name": ws.title.strip()}
    return None

def _parse_vessel_workbook(content: bytes):
    """Excel（.xlsx / .xlsm）から新造船リストを取り出す。MOL形式か常石形式かを自動判別する。
    戻り値: (source, vessels, info)"""
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise HTTPException(status_code=500, detail="サーバーに openpyxl がインストールされていません")
    try:
        wb = load_workbook(io.BytesIO(content), data_only=True)
    except Exception:
        raise HTTPException(status_code=400, detail="Excelファイル(.xlsx / .xlsm)として読み込めませんでした")
    for source, fn in (("mol", _parse_mol), ("tsuneishi", _parse_tsuneishi)):
        r = fn(wb)
        if r is not None:
            return source, r[0], r[1]
    raise HTTPException(
        status_code=400,
        detail="MOLの新造船工程表（VESSEL_NAME／PLANNED_DATE_OF_BUILD_DATE）か、常石のスケジュール表（番船／デリバリー日予定）の見出しが見つかりません",
    )

# -------------------------------------------------------
# 新造船データの保存先（機密）と閲覧パスワード
#   保存先: 非公開リポジトリ VESSEL_GH_REPO の vessels.json（GITHUB_TOKEN で読み書き）
#           GITHUB_TOKEN が無い環境(ローカル)では vessels.json をローカル保存（gitignore済み）
#   閲覧  : 環境変数 VESSEL_PASSWORD（チーム共通パスワード）。未設定なら全拒否(fail closed)
# -------------------------------------------------------
VESSEL_PASSWORD  = os.environ.get("VESSEL_PASSWORD", "")
VESSEL_GH_REPO   = os.environ.get("VESSEL_GH_REPO", "aihara1089-creator/monitoring-gr-koji-yotei-private")
VESSEL_GH_PATH   = os.environ.get("VESSEL_GH_PATH", "vessels.json")
VESSEL_GH_BRANCH = os.environ.get("VESSEL_GH_BRANCH", "main")
_VS_API          = f"https://api.github.com/repos/{VESSEL_GH_REPO}/contents/{VESSEL_GH_PATH}"
_VS_FILE         = os.path.join(_BASE_DIR, "vessels.json")
_VS_COOKIE       = "vessel_auth"
_VS_TTL          = 30 * 24 * 3600
_VS_LOCK         = threading.Lock()
_vs_sha          = ""
_vs_cache        = None   # {"vessels": [...], "sources": {"mol": {...}, "tsuneishi": {...}}}
_vs_fail         = {}     # ip -> [失敗時刻...]  （総当たり対策）
_vs_fail_all     = []     # 全体の失敗時刻

def _vs_empty():
    return {"vessels": [], "sources": {}}

def _vs_load() -> dict:
    """メモリ → GitHub(非公開) → ローカルファイル の順で取得"""
    global _vs_cache, _vs_sha
    if _vs_cache is not None:
        return _vs_cache
    data = None
    if GITHUB_TOKEN:
        try:
            req = urllib.request.Request(f"{_VS_API}?ref={VESSEL_GH_BRANCH}", headers=_gh_headers())
            with urllib.request.urlopen(req, timeout=10) as resp:
                meta = json.loads(resp.read())
            _vs_sha = meta.get("sha", "")
            data = json.loads(base64.b64decode(meta["content"]).decode("utf-8"))
        except urllib.error.HTTPError as e:
            if e.code != 404:
                print(f"[vessels fetch] HTTP {e.code}")
        except Exception as e:
            print(f"[vessels fetch] failed: {e}")
    if data is None and os.path.exists(_VS_FILE):
        try:
            with open(_VS_FILE, "r", encoding="utf-8") as f:
                data = json.load(f)
        except Exception:
            data = None
    _vs_cache = _vs_normalize(data) if isinstance(data, dict) else _vs_empty()
    return _vs_cache

def _vs_normalize(d: dict) -> dict:
    """旧形式（meta のみ・source なし）を sources 形式にそろえる"""
    d.setdefault("vessels", [])
    for v in d["vessels"]:
        v.setdefault("source", "mol")
    if "sources" not in d:
        d["sources"] = {"mol": d["meta"]} if d.get("meta") else {}
    d.pop("meta", None)
    return d

def _vs_push(data: dict):
    """非公開リポジトリに保存。失敗したら例外（取込を成功扱いにしない）"""
    global _vs_sha
    with open(_VS_FILE, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)
    if not GITHUB_TOKEN:
        return
    content_b64 = base64.b64encode(json.dumps(data, ensure_ascii=False, indent=2).encode("utf-8")).decode("ascii")
    last = None
    for attempt in range(3):
        payload = {"message": "chore: update vessels", "content": content_b64, "branch": VESSEL_GH_BRANCH}
        if _vs_sha:
            payload["sha"] = _vs_sha
        try:
            req = urllib.request.Request(
                _VS_API, data=json.dumps(payload).encode("utf-8"),
                headers={**_gh_headers(), "Content-Type": "application/json"}, method="PUT")
            with urllib.request.urlopen(req, timeout=20) as resp:
                _vs_sha = json.loads(resp.read()).get("content", {}).get("sha", _vs_sha)
            return
        except urllib.error.HTTPError as e:
            last = f"HTTP {e.code}"
            if e.code in (409, 422):      # SHA不一致 → 取り直す
                try:
                    r2 = urllib.request.Request(f"{_VS_API}?ref={VESSEL_GH_BRANCH}", headers=_gh_headers())
                    with urllib.request.urlopen(r2, timeout=10) as resp:
                        _vs_sha = json.loads(resp.read()).get("sha", "")
                except urllib.error.HTTPError:
                    _vs_sha = ""
                continue
            break
        except Exception as e:
            last = str(e)
            time.sleep(1)
    raise RuntimeError(f"非公開リポジトリへの保存に失敗しました ({last})")

def _vs_key() -> bytes:
    return hashlib.sha256(("vessel-auth|" + VESSEL_PASSWORD).encode("utf-8")).digest()

def _vs_make_token() -> str:
    exp = str(int(time.time()) + _VS_TTL)
    sig = hmac.new(_vs_key(), exp.encode(), hashlib.sha256).hexdigest()
    return f"{exp}.{sig}"

def _vs_token_ok(tok: str) -> bool:
    try:
        exp, sig = tok.split(".", 1)
        if int(exp) < time.time():
            return False
        good = hmac.new(_vs_key(), exp.encode(), hashlib.sha256).hexdigest()
        return hmac.compare_digest(sig, good)
    except Exception:
        return False

def _vs_require(request: Request):
    """パスワード未設定なら 503、未ログインなら 401"""
    if not VESSEL_PASSWORD:
        raise HTTPException(status_code=503, detail="船データの閲覧パスワードが未設定です（管理者が VESSEL_PASSWORD を設定してください）")
    if not _vs_token_ok(request.cookies.get(_VS_COOKIE, "")):
        raise HTTPException(status_code=401, detail="パスワードが必要です")

class VesselLogin(BaseModel):
    password: str

def _client_ip(request: Request) -> str:
    xf = request.headers.get("x-forwarded-for", "")
    return (xf.split(",")[0].strip() if xf else (request.client.host if request.client else "?"))

@app.post("/api/vessels/login")
def vessels_login(body: VesselLogin, request: Request):
    if not VESSEL_PASSWORD:
        raise HTTPException(status_code=503, detail="船データの閲覧パスワードが未設定です（管理者が VESSEL_PASSWORD を設定してください）")
    ip, now = _client_ip(request), time.time()
    fails = [t for t in _vs_fail.get(ip, []) if now - t < 600]
    _vs_fail_all[:] = [t for t in _vs_fail_all if now - t < 600]
    # IPごとの上限 + 全体の上限（IPを偽装した総当たりでも、短いパスワードを破られにくくする）
    if len(fails) >= 5 or len(_vs_fail_all) >= 20:
        raise HTTPException(status_code=429, detail="試行回数が多すぎます。10分ほど待ってからやり直してください")
    if not hmac.compare_digest(body.password.encode("utf-8"), VESSEL_PASSWORD.encode("utf-8")):
        fails.append(now)
        _vs_fail[ip] = fails
        _vs_fail_all.append(now)
        raise HTTPException(status_code=401, detail="パスワードが違います")
    _vs_fail.pop(ip, None)
    resp = JSONResponse({"ok": True})
    secure = request.headers.get("x-forwarded-proto", request.url.scheme) == "https"
    resp.set_cookie(_VS_COOKIE, _vs_make_token(), max_age=_VS_TTL, httponly=True,
                    samesite="lax", secure=secure, path="/api/vessels")
    return resp

@app.post("/api/vessels/logout")
def vessels_logout():
    resp = JSONResponse({"ok": True})
    resp.delete_cookie(_VS_COOKIE, path="/api/vessels")
    return resp

@app.get("/api/vessels")
def get_vessels(request: Request):
    _vs_require(request)
    with _VS_LOCK:
        d = _vs_load()
    resp = JSONResponse({"vessels": d.get("vessels", []), "sources": d.get("sources", {})})
    resp.headers["Cache-Control"] = "no-store"
    return resp

@app.post("/api/vessels/import")
async def import_vessels(request: Request, file: UploadFile = File(...)):
    _vs_require(request)
    filename = file.filename or ""
    if not filename.lower().endswith((".xlsx", ".xlsm")):
        raise HTTPException(status_code=400, detail="拡張子が .xlsx / .xlsm のファイルを選んでください")
    content = await file.read()
    if len(content) > _VESSEL_MAX_BYTES:
        raise HTTPException(status_code=413, detail="ファイルが大きすぎます（10MBまで）")
    source, vessels, info = await asyncio.to_thread(_parse_vessel_workbook, content)
    if not vessels:
        raise HTTPException(status_code=400, detail="船のデータが1件も見つかりませんでした")

    def _apply():
        global _vs_cache
        with _VS_LOCK:
            cur = _vs_load()
            # 差分比較・置き換えは同じ種類（MOL／常石）の中だけで行う
            same = [v for v in cur.get("vessels", []) if v.get("source", "mol") == source]
            others = [v for v in cur.get("vessels", []) if v.get("source", "mol") != source]
            old = {v.get("key"): v for v in same}
            changed = new = 0
            for v in vessels:
                o = old.get(v["key"])
                if o is None:
                    if old:                      # その種類の初回取込のときは「新規」扱いにしない
                        v["is_new"] = True
                        new += 1
                elif (o.get("delivery") or "") != (v["delivery"] or ""):
                    v["prev_delivery"] = o.get("delivery") or ""
                    changed += 1
            sources = dict(cur.get("sources", {}))
            sources[source] = {
                "source_file": filename,
                "sheet": info["sheet"],
                "table_name": info["table_name"],
                "imported_at": datetime.now(timezone.utc).isoformat(timespec="seconds"),
                "count": len(vessels),
            }
            data = {"vessels": others + vessels, "sources": sources}
            _vs_push(data)       # 保存に成功してからメモリへ反映
            _vs_cache = data
            return changed, new

    try:
        changed, new = await asyncio.to_thread(_apply)
    except RuntimeError as e:
        raise HTTPException(status_code=502, detail=str(e))

    dated = sum(1 for v in vessels if v["delivery_date"])
    return {
        "ok": True, "source": source, "label": _VESSEL_SOURCES[source],
        "count": len(vessels), "dated": dated, "tbd": len(vessels) - dated,
        "changed": changed, "new": new,
        "sheet": info["sheet"], "table_name": info["table_name"],
    }

# -------------------------------------------------------
# 静的ファイル配信
# -------------------------------------------------------
app.mount("/static", StaticFiles(directory=os.path.join(os.path.dirname(__file__), "static")), name="static")

@app.get("/")
def index():
    return FileResponse(os.path.join(os.path.dirname(__file__), "static", "index.html"))

@app.get("/{path:path}")
def catch_all(path: str):
    static_dir = os.path.join(os.path.dirname(__file__), "static")
    file_path = os.path.join(static_dir, path)
    if os.path.exists(file_path) and os.path.isfile(file_path):
        return FileResponse(file_path)
    return FileResponse(os.path.join(static_dir, "index.html"))

if __name__ == "__main__":
    import uvicorn
    uvicorn.run(app, host="0.0.0.0", port=8000)
